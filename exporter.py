# -*- coding: utf-8 -*-
"""将错题生成可编辑 Word 或适合打印的 PDF。"""
from __future__ import annotations

import html
import io
import re
import subprocess
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A3, A4, A5, B5, LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer
from PIL import Image, ImageDraw, ImageFont, ImageOps, UnidentifiedImageError
try:
    from latex2mathml.converter import convert as latex_to_mathml
    from mathml2omml import convert as mathml_to_omml
except ImportError:
    latex_to_mathml = None
    mathml_to_omml = None


PAPER_MM = {
    "A3": (297, 420), "A4": (210, 297), "A5": (148, 210),
    "B5": (176, 250), "Letter": (215.9, 279.4),
}
PDF_PAPERS = {"A3": A3, "A4": A4, "A5": A5, "B5": B5, "Letter": LETTER}


class _TextExtractor(HTMLParser):
    breaks = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs):
        if tag in self.breaks:
            self.parts.append("\n")
        if tag == "li":
            self.parts.append("• ")

    def handle_endtag(self, tag: str):
        if tag in self.breaks:
            self.parts.append("\n")

    def handle_data(self, data: str):
        self.parts.append(data)


def plain_text(value: Any) -> str:
    if not value:
        return ""
    parser = _TextExtractor()
    try:
        parser.feed(str(value))
        parser.close()
        text = "".join(parser.parts)
    except Exception:
        text = re.sub(r"<[^>]+>", "", str(value))
    text = html.unescape(text).replace("\xa0", " ")
    return re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", text).strip()


def normalize_formula_text(value: Any) -> str:
    """统一 OCR/139 文本中的公式转义，避免 Word 中出现双反斜杠或控制字符。"""
    text = str(value or "")
    # 模型有时把 JSON 中的 LaTeX 再多转义一层，只折叠公式命令和定界符，
    # 不动 aligned 环境里的换行双反斜杠。
    text = re.sub(r"\\{2,}(?=(?:[A-Za-z]{2,}|[()[\]{}]))", r"\\", text)
    text = text.replace("\x08", "\\b").replace("\x0c", "\\f").replace("\x0b", "\\v")
    text = re.sub(r"(?<!\\)\$\$([^$]+)\$\$", r"\\[\1\\]", text, flags=re.S)
    # 单美元包裹的数学表达式统一为 KaTeX/MathJax 通用的行内分隔符；
    # 普通金额或英文文本中的美元不转换。
    text = re.sub(
        r"(?<!\\)\$([^$\n]+?)\$",
        lambda match: (r"\(" + match.group(1) + r"\)")
        if re.search(r"(?:\\[A-Za-z]+|[\\^_=<>]|\d\s*[+\-*/]\s*\d)", match.group(1)) else match.group(0),
        text,
    )
    return text


_FORMULA_RE = re.compile(
    r"(\\\[.*?\\\]|\\\(.*?\\\)|\$\$.*?\$\$|\$(?=[^$\n]*?(?:\\[A-Za-z]+|[\^_=]))[^$\n]+\$)",
    re.S,
)
_BARE_FORMULA_RE = re.compile(
    r"(?<![\w\\])(?:[A-Za-z](?:[_^](?:\{[^{}\n]+\}|[A-Za-z0-9]))+)(?![\w])"
)


def _formula_parts(text: str) -> list[tuple[str, str, bool]]:
    """拆分普通文字和四种常见公式定界符。"""
    parts: list[tuple[str, str, bool]] = []
    cursor = 0
    matches = list(_FORMULA_RE.finditer(text))
    occupied = [(match.start(), match.end()) for match in matches]
    for match in _BARE_FORMULA_RE.finditer(text):
        if not any(start < match.end() and match.start() < end for start, end in occupied):
            matches.append(match)
    matches.sort(key=lambda match: match.start())
    for match in matches:
        if match.start() > cursor:
            parts.append(("text", text[cursor:match.start()], False))
        raw = match.group(0)
        if raw.startswith("\\[") and raw.endswith("\\]"):
            formula, display = raw[2:-2], True
        elif raw.startswith("\\(") and raw.endswith("\\)"):
            formula, display = raw[2:-2], False
        elif raw.startswith("$$") and raw.endswith("$$"):
            formula, display = raw[2:-2], True
        elif raw.startswith("$") and raw.endswith("$"):
            formula, display = raw[1:-1], False
        else:
            formula, display = raw, False
        parts.append(("formula", formula.strip(), display))
        cursor = match.end()
    if cursor < len(text):
        parts.append(("text", text[cursor:], False))
    return parts or [("text", text, False)]


_OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_OMML_COMMANDS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ϵ", "varepsilon": "ε",
    "zeta": "ζ", "eta": "η", "theta": "θ", "vartheta": "ϑ", "iota": "ι", "kappa": "κ",
    "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "omicron": "ο", "pi": "π",
    "varpi": "ϖ", "rho": "ρ", "varrho": "ϱ", "sigma": "σ", "varsigma": "ς", "tau": "τ",
    "upsilon": "υ", "phi": "ϕ", "varphi": "φ", "chi": "χ", "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ", "Pi": "Π",
    "Sigma": "Σ", "Upsilon": "Υ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
    "infty": "∞", "partial": "∂", "nabla": "∇", "forall": "∀", "exists": "∃", "emptyset": "∅",
    "pm": "±", "mp": "∓", "times": "×", "cdot": "⋅", "div": "÷", "ast": "∗", "star": "⋆",
    "circ": "∘", "bullet": "∙", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "neq": "≠",
    "ne": "≠", "approx": "≈", "equiv": "≡", "sim": "∼", "propto": "∝", "in": "∈", "notin": "∉",
    "subset": "⊂", "supset": "⊃", "subseteq": "⊆", "supseteq": "⊇", "to": "→", "rightarrow": "→",
    "leftarrow": "←", "leftrightarrow": "↔", "Rightarrow": "⇒", "Leftarrow": "⇐", "Leftrightarrow": "⇔",
    "ldots": "…", "dots": "…", "cdots": "⋯", "vdots": "⋮", "ddots": "⋱", "angle": "∠",
}
_OMML_ESCAPES = {"%": "%", "_": "_", "#": "#", "$": "$", "&": "&", "{": "{", "}": "}", " ": " "}
_OMML_NARY = {
    "sum": "∑", "prod": "∏", "coprod": "∐", "int": "∫", "iint": "∬", "iiint": "∭",
    "oint": "∮", "bigcup": "⋃", "bigcap": "⋂", "bigvee": "⋁", "bigwedge": "⋀",
}
_OMML_ACCENTS = {"hat": "^", "widehat": "^", "bar": "¯", "overline": "¯", "vec": "→", "tilde": "~"}
_OMML_FUNCTIONS = {"sin", "cos", "tan", "cot", "sec", "csc", "arcsin", "arccos", "arctan", "log", "ln", "lim", "max", "min", "det", "gcd"}


def _omml_element(name: str):
    return OxmlElement(f"m:{name}")


def _omml_run(text: str):
    run = _omml_element("r")
    text_node = _omml_element("t")
    if text[:1].isspace() or text[-1:].isspace():
        text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    return run


def _omml_container(name: str, nodes: list):
    container = _omml_element(name)
    for node in nodes:
        container.append(node)
    return container


def _omml_fraction(numerator: list, denominator: list, bar: bool = True):
    fraction = _omml_element("f")
    if not bar:
        props = _omml_element("fPr")
        fraction_type = _omml_element("type")
        fraction_type.set(qn("m:val"), "noBar")
        props.append(fraction_type)
        fraction.append(props)
    fraction.append(_omml_container("num", numerator))
    fraction.append(_omml_container("den", denominator))
    return fraction


def _omml_radical(radicand: list, degree: list | None = None):
    radical = _omml_element("rad")
    if degree is None:
        props = _omml_element("radPr")
        hidden = _omml_element("degHide")
        hidden.set(qn("m:val"), "1")
        props.append(hidden)
        radical.append(props)
    else:
        radical.append(_omml_container("deg", degree))
    radical.append(_omml_container("e", radicand))
    return radical


def _omml_accent(nodes: list, character: str):
    accent = _omml_element("acc")
    props = _omml_element("accPr")
    chr_node = _omml_element("chr")
    chr_node.set(qn("m:val"), character)
    props.append(chr_node)
    accent.append(props)
    accent.append(_omml_container("e", nodes))
    return accent


def _omml_nary(character: str):
    nary = _omml_element("nary")
    props = _omml_element("naryPr")
    chr_node = _omml_element("chr")
    chr_node.set(qn("m:val"), character)
    props.append(chr_node)
    nary.append(props)
    nary.append(_omml_container("sub", []))
    nary.append(_omml_container("sup", []))
    nary.append(_omml_container("e", []))
    return nary


def _omml_delimiter(nodes: list, begin: str, end: str):
    delimiter = _omml_element("d")
    props = _omml_element("dPr")
    begin_node = _omml_element("begChr")
    begin_node.set(qn("m:val"), "" if begin == "." else begin)
    end_node = _omml_element("endChr")
    end_node.set(qn("m:val"), "" if end == "." else end)
    props.extend([begin_node, end_node])
    delimiter.append(props)
    delimiter.append(_omml_container("e", nodes))
    return delimiter


def _omml_matrix(source: str, environment: str):
    rows = re.split(r"\\\\|\\newline", source)
    matrix = _omml_element("m")
    for row in rows:
        matrix_row = _omml_element("mr")
        for cell in row.split("&"):
            matrix_row.append(_omml_container("e", _LatexOMMLParser(cell.strip()).parse()))
        matrix.append(matrix_row)
    if environment in {"pmatrix", "bmatrix", "Bmatrix", "vmatrix", "Vmatrix"}:
        delimiters = {"pmatrix": ("(", ")"), "bmatrix": ("[", "]"), "Bmatrix": ("{", "}"), "vmatrix": ("|", "|"), "Vmatrix": ("‖", "‖")}
        begin, end = delimiters[environment]
        return _omml_delimiter([matrix], begin, end)
    return matrix


class _LatexOMMLParser:
    def __init__(self, text: str):
        self.text = text
        self.index = 0

    def parse(self, stop: str | None = None) -> list:
        nodes = []
        while self.index < len(self.text):
            if stop and self.text[self.index] == stop:
                self.index += 1
                break
            if self.text.startswith(r"\\", self.index):
                self.index += 2
                if stop:
                    break
                nodes.append(_omml_run(" "))
                continue
            if self.text[self.index].isspace():
                self.index += 1
                continue
            base = self._atom()
            if base is None:
                continue
            sub = sup = None
            while self.index < len(self.text) and self.text[self.index] in "^_":
                marker = self.text[self.index]
                self.index += 1
                argument = self._argument()
                if marker == "^":
                    sup = argument
                else:
                    sub = argument
            if sub is not None or sup is not None:
                base = self._scripts(base, sub, sup)
            nodes.append(base)
        return nodes

    def _argument(self) -> list:
        while self.index < len(self.text) and self.text[self.index].isspace():
            self.index += 1
        if self.index >= len(self.text):
            return []
        if self.text[self.index] == "{":
            self.index += 1
            return self.parse("}")
        node = self._atom()
        return [node] if node is not None else []

    def _raw_group(self) -> str:
        while self.index < len(self.text) and self.text[self.index].isspace():
            self.index += 1
        if self.index >= len(self.text) or self.text[self.index] != "{":
            return ""
        start = self.index = self.index + 1
        depth = 1
        while self.index < len(self.text) and depth:
            char = self.text[self.index]
            if char == "{" and (self.index == 0 or self.text[self.index - 1] != "\\"):
                depth += 1
            elif char == "}" and (self.index == 0 or self.text[self.index - 1] != "\\"):
                depth -= 1
            self.index += 1
        return self.text[start:self.index - 1]

    def _environment(self, name: str) -> str:
        marker = rf"\end{{{name}}}"
        end = self.text.find(marker, self.index)
        if end < 0:
            source = self.text[self.index:]
            self.index = len(self.text)
            return source
        source = self.text[self.index:end]
        self.index = end + len(marker)
        return source

    def _scripts(self, base, sub: list | None, sup: list | None):
        if sub is not None and sup is not None:
            node = _omml_element("sSubSup")
            node.extend([_omml_container("e", [base]), _omml_container("sub", sub), _omml_container("sup", sup)])
            return node
        if sub is not None:
            node = _omml_element("sSub")
            node.extend([_omml_container("e", [base]), _omml_container("sub", sub)])
            return node
        node = _omml_element("sSup")
        node.extend([_omml_container("e", [base]), _omml_container("sup", sup or [])])
        return node

    def _atom(self):
        char = self.text[self.index]
        if char == "{":
            self.index += 1
            return _omml_container("e", self.parse("}"))
        if char == "\\":
            return self._command()
        self.index += 1
        return _omml_run(char)

    def _command(self):
        self.index += 1
        if self.index >= len(self.text):
            return _omml_run("\\")
        if not self.text[self.index].isalpha():
            escaped = self.text[self.index]
            self.index += 1
            return _omml_run(_OMML_ESCAPES.get(escaped, escaped))
        start = self.index
        while self.index < len(self.text) and self.text[self.index].isalpha():
            self.index += 1
        command = self.text[start:self.index]
        if command in {"left", "right"}:
            delimiter = self._delimiter()
            return _omml_run("" if delimiter == "." else delimiter)
        if command in {"frac", "dfrac", "tfrac"}:
            return _omml_fraction(self._argument(), self._argument())
        if command in {"binom", "dbinom"}:
            return _omml_fraction(self._argument(), self._argument(), bar=False)
        if command == "sqrt":
            degree = None
            if self.index < len(self.text) and self.text[self.index] == "[":
                self.index += 1
                degree = self.parse("]")
            return _omml_radical(self._argument(), degree)
        if command in _OMML_NARY:
            return _omml_nary(_OMML_NARY[command])
        if command in _OMML_ACCENTS:
            return _omml_accent(self._argument(), _OMML_ACCENTS[command])
        if command in {"text", "textnormal", "textrm", "mathrm", "mathbf", "mathit", "operatorname"}:
            return _omml_run(self._raw_group())
        if command in {"left", "right"}:
            return _omml_run("")
        if command == "begin":
            environment = self._raw_group()
            return _omml_matrix(self._environment(environment), environment)
        if command in {"displaystyle", "textstyle", "scriptstyle", "scriptscriptstyle", "limits", "nolimits", "quad", "qquad", "!"}:
            return _omml_run(" ")
        if command in _OMML_FUNCTIONS:
            return _omml_run(command)
        return _omml_run(_OMML_COMMANDS.get(command, f"\\{command}"))

    def _delimiter(self) -> str:
        while self.index < len(self.text) and self.text[self.index].isspace():
            self.index += 1
        if self.index >= len(self.text):
            return "."
        if self.text[self.index] == "\\":
            command = self._command()
            text = "".join(node.findtext(f"{{{_OMML_NS}}}t") or "" for node in [command])
            return text or "."
        delimiter = self.text[self.index]
        self.index += 1
        return delimiter


def _omml_formula(formula: str, display: bool = False):
    normalized = normalize_formula_text(formula).strip()
    omml = None
    if latex_to_mathml is not None and mathml_to_omml is not None:
        try:
            mathml = latex_to_mathml(normalized, display="block" if display else "inline")
            omml_text = mathml_to_omml(mathml)
            omml_text = omml_text.replace("<m:oMath>", f'<m:oMath xmlns:m="{_OMML_NS}">', 1)
            omml = parse_xml(omml_text.encode("utf-8"))
        except Exception:
            omml = None
    if omml is None:
        formula_root = _omml_element("oMath")
        formula_root.append(_omml_container("e", _LatexOMMLParser(normalized).parse()))
        omml = formula_root
    if display:
        paragraph = _omml_element("oMathPara")
        paragraph.append(omml)
        return paragraph
    return omml


def _append_formula_paragraph(document: Document, text: str, usable_width: float, font_size: float) -> None:
    """添加可混排的文字与 Word 原生公式段落。"""
    for line in normalize_formula_text(text).splitlines() or [""]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True
        parts = _formula_parts(line)
        has_display = any(kind == "formula" and display for kind, _value, display in parts)
        if has_display:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, value, display in parts:
            if kind == "text":
                if value:
                    paragraph.add_run(value)
                continue
            try:
                paragraph._p.append(_omml_formula(value, display))
            except Exception:
                paragraph.add_run(value)


def _options(raw: dict[str, Any] | None) -> dict[str, Any]:
    raw = raw or {}
    return {
        "title": str(raw.get("title") or "我的错题集"),
        "subtitle": str(raw.get("subtitle") or "系统整理 · 针对复习"),
        "paper": raw.get("paper") if raw.get("paper") in PAPER_MM else "A4",
        "orientation": "landscape" if raw.get("orientation") == "landscape" else "portrait",
        "margin_top": float(raw.get("margin_top", 18)),
        "margin_bottom": float(raw.get("margin_bottom", 18)),
        "margin_left": float(raw.get("margin_left", 18)),
        "margin_right": float(raw.get("margin_right", 18)),
        "line_spacing": max(1.0, min(3.0, float(raw.get("line_spacing", 1.5)))),
        "font_size": max(8, min(24, float(raw.get("font_size", 11)))),
        "blank_lines": max(0, min(12, int(raw.get("blank_lines", 3)))),
        "answer_mode": raw.get("answer_mode") if raw.get("answer_mode") in {"inline", "separate", "none"} else "separate",
        "include_analysis": bool(raw.get("include_analysis", True)),
        "include_my_answer": bool(raw.get("include_my_answer", True)),
        "include_note": bool(raw.get("include_note", True)),
        "include_meta": bool(raw.get("include_meta", True)),
        "page_break_each": bool(raw.get("page_break_each", False)),
        "page_numbers": bool(raw.get("page_numbers", True)),
    }


class _RichContentParser(HTMLParser):
    """把 139 返回的 HTML 转为按顺序排列的文字/图片事件。"""

    block_tags = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "section"}

    def __init__(self):
        super().__init__()
        self.events: list[tuple[str, str]] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs):
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "iframe", "object", "embed"}:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        if tag == "img":
            src = dict(attrs).get("src") or ""
            if src:
                self.events.append(("image", html.unescape(src)))
        elif tag == "li":
            self.events.append(("text", "• "))
        if tag in self.block_tags:
            self.events.append(("break", ""))

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "iframe", "object", "embed"}:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if not self.skip_depth and tag in self.block_tags:
            self.events.append(("break", ""))

    def handle_data(self, data: str):
        if not self.skip_depth and data:
            self.events.append(("text", html.unescape(data)))


def rich_events(value: Any) -> list[tuple[str, str]]:
    if not value:
        return []
    parser = _RichContentParser()
    try:
        parser.feed(str(value))
        parser.close()
        return parser.events
    except Exception:
        return [("text", plain_text(value))]


def remote_question_image_url(src: str) -> str:
    """还原 139 图片地址；明确拒绝本机上传图和 data URL。"""
    src = str(src or "").strip()
    if src.startswith("/proxy?url="):
        from urllib.parse import parse_qs, urlsplit
        src = parse_qs(urlsplit(src).query).get("url", [""])[0]
    elif src.startswith("//"):
        src = "https:" + src
    return src if src.startswith(("http://", "https://")) else ""


def question_image_sources(value: Any) -> list[str]:
    """返回题干中可从 139 获取的图片地址，本地图片和 data URL 会被忽略。"""
    return [url for kind, src in rich_events(value) if kind == "image" and (url := remote_question_image_url(src))]


def _layout_font_path() -> str:
    candidates = [
        Path(__file__).parent / "assets" / "NotoSansSC-Regular.ttf",
        Path("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    ]
    return str(next((path for path in candidates if path.exists()), ""))


def _layout_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = _layout_font_path()
    if path:
        try:
            return ImageFont.truetype(path, max(8, int(size)))
        except Exception:
            pass
    return ImageFont.load_default()


class _PagedImage:
    """固定纸张尺寸分页排版，最后把所有页面纵向合成为一张 PNG。"""

    def __init__(self, width: int, height: int, margins: tuple[int, int, int, int], line_spacing: float):
        self.width, self.height = width, height
        self.left, self.top, self.right, self.bottom = margins
        self.line_spacing = line_spacing
        self.pages: list[Image.Image] = []
        self._new_page()

    @property
    def draw(self) -> ImageDraw.ImageDraw:
        return self._draw

    @property
    def max_width(self) -> int:
        return self.width - self.left - self.right

    def _new_page(self) -> None:
        page = Image.new("RGB", (self.width, self.height), "white")
        self.pages.append(page)
        self._page = page
        self._draw = ImageDraw.Draw(page)
        self.y = self.top

    def page_break(self) -> None:
        self._new_page()

    def _line_height(self, font: ImageFont.ImageFont) -> int:
        box = self._draw.textbbox((0, 0), "中文Ag", font=font)
        return max(14, box[3] - box[1])

    def _ensure(self, height: int) -> None:
        if self.y > self.top and self.y + height > self.height - self.bottom:
            self._new_page()

    def _wrap(self, text: str, font: ImageFont.ImageFont) -> list[str]:
        lines: list[str] = []
        for paragraph in str(text).replace("\r", "").split("\n"):
            if not paragraph:
                lines.append("")
                continue
            current = ""
            for char in paragraph:
                candidate = current + char
                if current and self._draw.textlength(candidate, font=font) > self.max_width:
                    lines.append(current)
                    current = char
                else:
                    current = candidate
            lines.append(current)
        return lines or [""]

    def text(self, value: Any, font: ImageFont.ImageFont, fill: str = "#263142", after: int = 7,
             align: str = "left") -> None:
        lines = self._wrap(str(value or ""), font)
        line_height = int(self._line_height(font) * self.line_spacing)
        for line in lines:
            self._ensure(line_height)
            box = self._draw.textbbox((0, 0), line, font=font)
            text_width = box[2] - box[0]
            x = self.left if align == "left" else max(self.left, (self.width - text_width) // 2)
            self._draw.text((x, self.y), line, font=font, fill=fill)
            self.y += line_height
        self.y += after

    def rule(self, color: str = "#DDE2E7", gap: int = 13) -> None:
        self._ensure(gap + 2)
        self._draw.line((self.left, self.y, self.width - self.right, self.y), fill=color, width=2)
        self.y += gap

    def blank_lines(self, count: int, font: ImageFont.ImageFont) -> None:
        line_height = int(self._line_height(font) * self.line_spacing * 1.15)
        for _ in range(max(0, count)):
            self._ensure(line_height)
            self._draw.line((self.left, self.y + line_height // 2, self.width - self.right, self.y + line_height // 2), fill="#AEB7C2", width=1)
            self.y += line_height
        self.y += 5

    def image(self, image: Image.Image) -> None:
        image = image.convert("RGB")
        max_width = self.max_width
        max_height = max(180, self.height - self.top - self.bottom - 30)
        scale = min(1.0, max_width / image.width, max_height / image.height)
        width, height = max(1, int(image.width * scale)), max(1, int(image.height * scale))
        if self.y + height > self.height - self.bottom and self.y > self.top:
            self._new_page()
        if scale < 1:
            image = image.resize((width, height), Image.Resampling.LANCZOS)
        self._page.paste(image, (self.left + (max_width - width) // 2, self.y))
        self.y += height + 12

    def finish(self) -> bytes:
        gap = 18
        if len(self.pages) > 30:
            raise ValueError("OCR Word 单次排版超过 30 页，请缩小导出范围")
        output = Image.new("RGB", (self.width, self.height * len(self.pages) + gap * max(0, len(self.pages) - 1)), "white")
        y = 0
        for index, page in enumerate(self.pages):
            output.paste(page, (0, y))
            y += self.height + gap
        result = io.BytesIO()
        output.save(result, format="PNG", optimize=True, dpi=(150, 150))
        return result.getvalue()


def _build_layout_image_legacy(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    """把错题排成连续长图，供 CamScanner 图片转 Word OCR。"""
    opt = _options(raw_options)
    dpi = 150
    width_mm, height_mm = PAPER_MM[opt["paper"]]
    if opt["orientation"] == "landscape":
        width_mm, height_mm = height_mm, width_mm
    px = lambda mm_value: max(1, round(mm_value / 25.4 * dpi))
    width, height = px(width_mm), px(height_mm)
    layout = _PagedImage(width, height, (px(opt["margin_left"]), px(opt["margin_top"]), px(opt["margin_right"]), px(opt["margin_bottom"])), opt["line_spacing"])
    body = _layout_font(opt["font_size"] * dpi / 72)
    h1 = _layout_font(opt["font_size"] * 1.45 * dpi / 72)
    h2 = _layout_font(opt["font_size"] * 1.12 * dpi / 72)
    title_font = _layout_font(opt["font_size"] * 2.0 * dpi / 72)
    meta_font = _layout_font(opt["font_size"] * .78 * dpi / 72)
    layout.text(opt["title"], title_font, fill="#167A78", after=5, align="center")
    if opt["subtitle"]:
        layout.text(opt["subtitle"], meta_font, fill="#77808E", after=3, align="center")
    layout.text(f"共 {len(items)} 道错题", meta_font, fill="#77808E", after=20, align="center")

    for index, item in enumerate(items, 1):
        layout.text(f"{index}. {plain_text(item.get('title')) or '错题'}", h1, fill="#167A78", after=4)
        if opt["include_meta"]:
            tags = " · ".join(item.get("tags") or [])
            meta = "  |  ".join(filter(None, [item.get("subject"), item.get("grade"), item.get("question_type"), f"难度 {item.get('difficulty', 3)}/5", tags]))
            if meta:
                layout.text(meta, meta_font, fill="#77808E", after=8)
        layout.text("题目", h2, fill="#167A78", after=3)
        image_iter = iter(item.get("_question_images") or [])
        pending = ""
        rendered = False
        for kind, value in rich_events(item.get("question")):
            if kind == "text":
                pending += value
            elif kind == "break":
                if pending.strip():
                    layout.text(pending.strip(), body, after=2)
                    rendered = True
                pending = ""
            elif kind == "image":
                if pending.strip():
                    layout.text(pending.strip(), body, after=2)
                    rendered = True
                pending = ""
                if not remote_question_image_url(value):
                    continue
                try:
                    remote_image = next(image_iter)
                    if remote_image is not None:
                        layout.image(remote_image)
                        rendered = True
                except (StopIteration, UnidentifiedImageError, OSError, ValueError):
                    pass
        if pending.strip():
            layout.text(pending.strip(), body, after=5)
            rendered = True
        if not rendered:
            layout.text("（139 未返回可用题干）", body, fill="#77808E", after=5)
        layout.blank_lines(opt["blank_lines"], body)
        if opt["include_my_answer"] and plain_text(item.get("my_answer")):
            layout.text("原作答", h2, fill="#167A78", after=3)
            layout.text(plain_text(item.get("my_answer")), body, after=5)
        if opt["answer_mode"] == "inline":
            layout.text("答案", h2, fill="#167A78", after=3)
            layout.text(plain_text(item.get("answer")) or "（暂无）", body, after=4)
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                layout.text("解析", h2, fill="#167A78", after=3)
                layout.text(plain_text(item.get("analysis")), body, after=5)
        if opt["include_note"] and plain_text(item.get("note")):
            layout.text("复盘笔记", h2, fill="#167A78", after=3)
            layout.text(plain_text(item.get("note")), body, after=5)
        if opt["page_break_each"] and index != len(items):
            layout.page_break()
        elif index != len(items):
            layout.rule()

    if opt["answer_mode"] == "separate":
        layout.page_break()
        layout.text("参考答案与解析", title_font, fill="#167A78", after=18, align="center")
        for index, item in enumerate(items, 1):
            layout.text(f"第 {index} 题", h1, fill="#167A78", after=4)
            layout.text(plain_text(item.get("answer")) or "（暂无答案）", body, after=4)
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                layout.text("解析", h2, fill="#167A78", after=3)
                layout.text(plain_text(item.get("analysis")), body, after=5)
    return layout.finish()


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _word_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def build_docx(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    opt = _options(raw_options)
    document = Document()
    section = document.sections[0]
    width, height = PAPER_MM[opt["paper"]]
    if opt["orientation"] == "landscape":
        width, height = height, width
        section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(width / 10), Cm(height / 10)
    section.top_margin, section.bottom_margin = Cm(opt["margin_top"] / 10), Cm(opt["margin_bottom"] / 10)
    section.left_margin, section.right_margin = Cm(opt["margin_left"] / 10), Cm(opt["margin_right"] / 10)
    styles = document.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(opt["font_size"])
    styles["Normal"].paragraph_format.line_spacing = opt["line_spacing"]
    styles["Normal"].paragraph_format.space_after = Pt(5)
    title = document.add_heading(opt["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(opt["subtitle"], style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary = document.add_paragraph(f"共 {len(items)} 道错题")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    usable_width = max(6.0, width - opt["margin_left"] - opt["margin_right"])
    for index, item in enumerate(items, 1):
        heading = document.add_heading(f"{index}. {plain_text(item.get('title')) or '错题'}", level=1)
        heading.paragraph_format.keep_with_next = True
        if opt["include_meta"]:
            tags = " · ".join(item.get("tags") or [])
            meta = "  |  ".join(filter(None, [item.get("subject"), item.get("grade"), item.get("question_type"), f"难度 {item.get('difficulty', 3)}/5", tags]))
            if meta:
                p = document.add_paragraph(meta)
                p.style = styles["Subtitle"]
        document.add_heading("题目", level=2)
        question = plain_text(item.get("question"))
        if question:
            _append_formula_paragraph(document, question, usable_width, opt["font_size"])
        else:
            document.add_paragraph("（题目图片见原记录）")
        for _ in range(opt["blank_lines"]):
            document.add_paragraph("________________________________________________________________")
        if opt["include_my_answer"] and plain_text(item.get("my_answer")):
            document.add_heading("原作答", level=2)
            _append_formula_paragraph(document, plain_text(item.get("my_answer")), usable_width, opt["font_size"])
        if opt["answer_mode"] == "inline":
            document.add_heading("答案", level=2)
            answer = plain_text(item.get("answer"))
            _append_formula_paragraph(document, answer or "（暂无）", usable_width, opt["font_size"])
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                document.add_heading("解析", level=2)
                _append_formula_paragraph(document, plain_text(item.get("analysis")), usable_width, opt["font_size"])
        if opt["include_note"] and plain_text(item.get("note")):
            document.add_heading("复盘笔记", level=2)
            _append_formula_paragraph(document, plain_text(item.get("note")), usable_width, opt["font_size"])
        if opt["page_break_each"] and index != len(items):
            document.add_page_break()
        elif index != len(items):
            document.add_paragraph("—" * 24)
    if opt["answer_mode"] == "separate":
        document.add_page_break()
        document.add_heading("参考答案与解析", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        for index, item in enumerate(items, 1):
            document.add_heading(f"第 {index} 题", level=1)
            answer = plain_text(item.get("answer"))
            _append_formula_paragraph(document, answer or "（暂无答案）", usable_width, opt["font_size"])
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                document.add_heading("解析", level=2)
                _append_formula_paragraph(document, plain_text(item.get("analysis")), usable_width, opt["font_size"])
    if opt["page_numbers"]:
        _word_page_number(section.footer.paragraphs[0])
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_ai_docx(pages: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    """把视觉模型识别出的文字块和裁切图片块排成可编辑 Word。"""
    opt = _options(raw_options)
    document = Document()
    section = document.sections[0]
    width, height = PAPER_MM[opt["paper"]]
    if opt["orientation"] == "landscape":
        width, height = height, width
        section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(width / 10), Cm(height / 10)
    section.top_margin, section.bottom_margin = Cm(opt["margin_top"] / 10), Cm(opt["margin_bottom"] / 10)
    section.left_margin, section.right_margin = Cm(opt["margin_left"] / 10), Cm(opt["margin_right"] / 10)
    styles = document.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(opt["font_size"])
    styles["Normal"].paragraph_format.line_spacing = opt["line_spacing"]
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(25)
    styles["Subtitle"].font.size = Pt(9)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)

    title = document.add_heading(opt["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(opt["subtitle"], style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary = document.add_paragraph(f"混合排版 {len(pages)} 个内容单元 · 直接文字保留 · 图片按内容裁切")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()

    usable_width = max(60.0, width - opt["margin_left"] - opt["margin_right"])
    usable_height = max(80.0, height - opt["margin_top"] - opt["margin_bottom"])
    max_image_width = max(8.0, min(17.2, usable_width / 10 - 0.4))
    max_image_height = max(8.0, min(23.5, usable_height / 10 - 1.0))

    def configure_paragraph(paragraph, before=0, after=0, keep=False) -> None:
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.keep_together = keep
        paragraph.paragraph_format.widow_control = True

    def render_blocks(blocks: list[dict[str, Any]]) -> None:
        for block in blocks:
            kind = str(block.get("type") or "text").lower()
            if kind == "image" and block.get("data"):
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                configure_paragraph(paragraph, before=5, after=2, keep=True)
                stream = io.BytesIO(block["data"])
                try:
                    with Image.open(io.BytesIO(block["data"])) as source:
                        image = ImageOps.exif_transpose(source)
                        image_width, image_height = image.size
                    if not image_width or not image_height:
                        continue
                    scale = min(
                        max_image_width / (image_width / 96 * 2.54),
                        max_image_height / (image_height / 96 * 2.54),
                        1.0,
                    )
                    picture_width = image_width / 96 * 2.54 * scale
                    picture_height = image_height / 96 * 2.54 * scale
                    paragraph.add_run().add_picture(
                        stream, width=Cm(picture_width), height=Cm(picture_height)
                    )
                except (ValueError, OSError, UnidentifiedImageError):
                    continue
                caption = str(block.get("caption") or "").strip()
                if caption:
                    caption_paragraph = document.add_paragraph(caption)
                    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_paragraph.style = styles["Subtitle"]
                    configure_paragraph(caption_paragraph, after=5, keep=True)
                continue
            if kind in {"image-pending", "image"}:
                continue
            text = normalize_formula_text(block.get("text"))
            text = text.strip()
            if not text:
                continue
            style = str(block.get("style") or "text")
            if style == "title":
                paragraph = document.add_heading(text, level=1)
                configure_paragraph(paragraph, before=4, after=6, keep=True)
            elif style == "heading":
                paragraph = document.add_heading(text, level=2)
                configure_paragraph(paragraph, before=9, after=3, keep=True)
            elif style == "meta":
                paragraph = document.add_paragraph(text, style="Subtitle")
                configure_paragraph(paragraph, after=8, keep=True)
            else:
                _append_formula_paragraph(document, text, usable_width, opt["font_size"])
                continue

    for page_index, page in enumerate(pages, 1):
        if page_index > 1:
            document.add_page_break()
        if not page.get("mixed"):
            heading = document.add_heading(f"第 {page_index} 页", level=1)
            heading.paragraph_format.keep_with_next = True
        blocks = page.get("blocks") or []
        render_blocks(blocks)
        if not blocks:
            document.add_paragraph("（AI 未识别到可用内容）")

    if opt["answer_mode"] == "separate":
        answer_pages = [page.get("answer_blocks") or [] for page in pages]
        if any(answer_pages):
            document.add_page_break()
            heading = document.add_heading("参考答案与解析", 0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for page_index, blocks in enumerate(answer_pages, 1):
                if page_index > 1:
                    document.add_page_break()
                document.add_heading(f"第 {page_index} 题", level=1)
                render_blocks(blocks)

    if opt["page_numbers"]:
        _word_page_number(section.footer.paragraphs[0])
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_visual_docx(layout_png: bytes, raw_options: dict[str, Any] | None) -> bytes:
    """CamScanner 不可用时，把已排好的每一页原样放进 Word，避免公式和配图丢失。"""
    opt = _options(raw_options)
    width_mm, height_mm = PAPER_MM[opt["paper"]]
    if opt["orientation"] == "landscape":
        width_mm, height_mm = height_mm, width_mm
    dpi = 150
    page_width = round(width_mm / 25.4 * dpi)
    page_height = round(height_mm / 25.4 * dpi)
    gap = 18
    with Image.open(io.BytesIO(layout_png)) as source:
        combined = source.convert("RGB")
    page_count = max(1, round((combined.height + gap) / (page_height + gap)))
    if combined.width != page_width or page_count > 30:
        raise ValueError("版式图片尺寸异常")

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(width_mm / 10), Cm(height_mm / 10)
    section.top_margin = section.bottom_margin = Cm(0.2)
    section.left_margin = section.right_margin = Cm(0.2)
    usable_width = Cm((width_mm - 4) / 10)
    for index in range(page_count):
        top = index * (page_height + gap)
        page = combined.crop((0, top, page_width, min(top + page_height, combined.height)))
        stream = io.BytesIO()
        page.save(stream, format="PNG", optimize=True)
        stream.seek(0)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.page_break_before = index > 0
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(stream, width=usable_width)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_font() -> str:
    candidates = [
        Path(__file__).parent / "assets" / "NotoSansSC-Regular.ttf",
        Path("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    ]
    for path in candidates:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont("MistakeCN", str(path)))
                return "MistakeCN"
            except Exception:
                continue
    return "Helvetica"


def _pdf_safe(text: Any) -> str:
    escaped = html.escape(plain_text(text))
    return escaped.replace("\n", "<br/>") or "（暂无）"


def _build_pdf_legacy(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    opt = _options(raw_options)
    paper = PDF_PAPERS[opt["paper"]]
    if opt["orientation"] == "landscape":
        paper = landscape(paper)
    output = io.BytesIO()
    font = _pdf_font()
    doc = SimpleDocTemplate(
        output, pagesize=paper,
        topMargin=opt["margin_top"] * mm, bottomMargin=opt["margin_bottom"] * mm,
        leftMargin=opt["margin_left"] * mm, rightMargin=opt["margin_right"] * mm,
        title=opt["title"], author="知错 · AI 错题本",
    )
    samples = getSampleStyleSheet()
    body = ParagraphStyle(
        "CNBody", parent=samples["BodyText"], fontName=font,
        fontSize=opt["font_size"], leading=opt["font_size"] * opt["line_spacing"],
        textColor=colors.HexColor("#263142"), spaceAfter=5 * mm,
    )
    h1 = ParagraphStyle("CNH1", parent=body, fontSize=18, leading=24, textColor=colors.HexColor("#167A78"), spaceBefore=5 * mm, spaceAfter=3 * mm)
    h2 = ParagraphStyle("CNH2", parent=body, fontSize=11, leading=16, textColor=colors.HexColor("#167A78"), spaceBefore=3 * mm, spaceAfter=1.5 * mm)
    meta_style = ParagraphStyle("CNMeta", parent=body, fontSize=8.5, leading=12, textColor=colors.HexColor("#77808E"), spaceAfter=3 * mm)
    cover = ParagraphStyle("CNCover", parent=h1, fontSize=26, leading=34, alignment=TA_CENTER, spaceAfter=6 * mm)
    centered = ParagraphStyle("CNCenter", parent=meta_style, alignment=TA_CENTER)
    story = [Paragraph(_pdf_safe(opt["title"]), cover), Paragraph(_pdf_safe(opt["subtitle"]), centered), Paragraph(f"共 {len(items)} 道错题", centered), Spacer(1, 8 * mm)]
    for index, item in enumerate(items, 1):
        story.append(Paragraph(f"{index}. {_pdf_safe(item.get('title') or '错题')}", h1))
        if opt["include_meta"]:
            tags = " · ".join(item.get("tags") or [])
            meta = "  |  ".join(filter(None, [item.get("subject"), item.get("grade"), item.get("question_type"), f"难度 {item.get('difficulty', 3)}/5", tags]))
            if meta:
                story.append(Paragraph(_pdf_safe(meta), meta_style))
        story.extend([Paragraph("题目", h2), Paragraph(_pdf_safe(item.get("question")) or "（题目图片见原记录）", body)])
        for _ in range(opt["blank_lines"]):
            story.extend([HRFlowable(width="100%", thickness=.35, color=colors.HexColor("#C7CDD4")), Spacer(1, 5 * mm)])
        if opt["include_my_answer"] and plain_text(item.get("my_answer")):
            story.extend([Paragraph("原作答", h2), Paragraph(_pdf_safe(item.get("my_answer")), body)])
        if opt["answer_mode"] == "inline":
            story.extend([Paragraph("答案", h2), Paragraph(_pdf_safe(item.get("answer")), body)])
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                story.extend([Paragraph("解析", h2), Paragraph(_pdf_safe(item.get("analysis")), body)])
        if opt["include_note"] and plain_text(item.get("note")):
            story.extend([Paragraph("复盘笔记", h2), Paragraph(_pdf_safe(item.get("note")), body)])
        if opt["page_break_each"] and index != len(items):
            story.append(PageBreak())
        elif index != len(items):
            story.extend([Spacer(1, 3 * mm), HRFlowable(width="100%", thickness=.7, color=colors.HexColor("#DDE2E7")), Spacer(1, 3 * mm)])
    if opt["answer_mode"] == "separate":
        story.extend([PageBreak(), Paragraph("参考答案与解析", cover)])
        for index, item in enumerate(items, 1):
            story.extend([Paragraph(f"第 {index} 题", h1), Paragraph(_pdf_safe(item.get("answer")), body)])
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                story.extend([Paragraph("解析", h2), Paragraph(_pdf_safe(item.get("analysis")), body)])

    def footer(canvas, current_doc):
        if not opt["page_numbers"]:
            return
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#8A929E"))
        canvas.drawCentredString(paper[0] / 2, 8 * mm, f"— {current_doc.page} —")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _image_data_url(image: Image.Image) -> str:
    output = io.BytesIO()
    image.convert("RGB").save(output, format="JPEG", quality=92, optimize=True)
    import base64
    return "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode("ascii")


class _ExportHTMLParser(HTMLParser):
    """仅保留打印所需标签，并把已下载的 139 图片嵌入文档。"""

    allowed_tags = {
        "p", "div", "br", "ul", "ol", "li", "table", "thead", "tbody", "tr", "td", "th",
        "b", "strong", "i", "em", "u", "sub", "sup", "span",
    }
    skipped_tags = {"script", "style", "noscript", "iframe", "object", "embed", "link"}

    def __init__(self, images: dict[str, Image.Image | None]):
        super().__init__(convert_charrefs=True)
        self.images = images
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs):
        tag = tag.lower()
        if tag in self.skipped_tags:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        attributes = dict(attrs)
        if tag == "img":
            source = remote_question_image_url(attributes.get("src") or "")
            image = self.images.get(source)
            if image is not None:
                floated = "float:right" in (attributes.get("style") or "").replace(" ", "").lower()
                css_class = "remote-image float-right" if floated else "remote-image"
                self.parts.append(f'<img class="{css_class}" src="{_image_data_url(image)}" alt="题目配图">')
            return
        if tag not in self.allowed_tags:
            return
        if tag == "br":
            self.parts.append("<br>")
            return
        safe_attrs = ""
        if tag in {"td", "th"}:
            pairs = []
            for name in ("colspan", "rowspan"):
                value = attributes.get(name, "")
                if str(value).isdigit():
                    pairs.append(f'{name}="{int(value)}"')
            safe_attrs = (" " + " ".join(pairs)) if pairs else ""
        self.parts.append(f"<{tag}{safe_attrs}>")

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        if tag in self.skipped_tags:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if not self.skip_depth and tag in self.allowed_tags and tag != "br":
            self.parts.append(f"</{tag}>")

    def handle_data(self, data: str):
        if not self.skip_depth:
            self.parts.append(html.escape(data, quote=False))


def _rich_export_html(value: Any, images: dict[str, Image.Image | None]) -> str:
    if not value:
        return ""
    parser = _ExportHTMLParser(images)
    try:
        parser.feed(str(value))
        parser.close()
        return "".join(parser.parts)
    except Exception:
        return html.escape(plain_text(value)).replace("\n", "<br>")


def _export_html(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> str:
    opt = _options(raw_options)
    width_mm, height_mm = PAPER_MM[opt["paper"]]
    if opt["orientation"] == "landscape":
        width_mm, height_mm = height_mm, width_mm
    page_number_css = ""
    if opt["page_numbers"]:
        page_number_css = '@bottom-center { content: "第 " counter(page) " 页"; color: #7b858e; font-size: 8pt; }'
    css = f"""
      @page {{ size: {width_mm}mm {height_mm}mm; margin: {opt['margin_top']}mm {opt['margin_right']}mm {opt['margin_bottom']}mm {opt['margin_left']}mm; {page_number_css} }}
      * {{ box-sizing: border-box; }}
      html, body {{ margin: 0; padding: 0; }}
      body {{ color: #242c33; font-family: "Droid Sans Fallback", "WenQuanYi Zen Hei", Arial, sans-serif; font-size: {opt['font_size']}pt; line-height: {opt['line_spacing']}; print-color-adjust: exact; -webkit-print-color-adjust: exact; }}
      h1, h2, p {{ margin-top: 0; }}
      .document-title {{ margin: 0 0 2mm; color: #166f6c; font-size: {opt['font_size'] * 2.0}pt; line-height: 1.3; text-align: center; letter-spacing: 0; }}
      .subtitle, .summary {{ margin: 0 0 2mm; color: #737e87; font-size: {max(8, opt['font_size'] * .78)}pt; text-align: center; }}
      .summary {{ margin-bottom: 8mm; }}
      .question-block {{ clear: both; margin: 0 0 7mm; }}
      .question-heading {{ break-inside: avoid; break-after: avoid; }}
      .question-title {{ margin: 0 0 1.5mm; color: #166f6c; font-size: {opt['font_size'] * 1.28}pt; line-height: 1.45; break-after: avoid; }}
      .meta {{ margin: 0 0 2.5mm; color: #77818a; font-size: {max(8, opt['font_size'] * .76)}pt; }}
      .section-title {{ margin: 3mm 0 1.5mm; color: #166f6c; font-size: {opt['font_size'] * 1.03}pt; break-after: avoid; }}
      .content {{ overflow-wrap: anywhere; }}
      .content::after {{ display: block; clear: both; content: ""; }}
      .content p, .content div {{ margin: 0 0 1.7mm; }}
      .content ul, .content ol {{ margin: 1.4mm 0 2mm; padding-left: 7mm; }}
      .content li {{ margin: .7mm 0; }}
      .content table {{ width: 100%; margin: 2mm 0; border-collapse: collapse; }}
      .content td, .content th {{ padding: 1.3mm 1.8mm; border: .25mm solid #bfc7cc; }}
      .remote-image {{ display: block; max-width: 72%; max-height: 78mm; width: auto; height: auto; margin: 2mm auto 3mm; object-fit: contain; }}
      .remote-image.float-right {{ float: right; max-width: 45%; margin: 0 0 3mm 5mm; }}
      .katex {{ font-size: 1.04em; }}
      .katex-display {{ margin: 2.5mm 0; overflow: hidden; }}
      .answer-line {{ height: 8mm; border-bottom: .25mm solid #c9cfd3; }}
      .answer-box {{ margin-top: 2mm; padding: 2.5mm 3mm; border-left: 1mm solid #27847f; background: #eef6f4; }}
      .divider {{ clear: both; height: .25mm; margin: 6mm 0; background: #d8dddf; }}
      .page-break {{ break-before: page; }}
      .force-page {{ break-after: page; }}
    """
    sections: list[str] = []
    for index, item in enumerate(items, 1):
        images = item.get("_remote_images") or {}
        tags = " · ".join(item.get("tags") or [])
        meta = "  |  ".join(filter(None, [item.get("subject"), item.get("grade"), item.get("question_type"), f"难度 {item.get('difficulty', 3)}/5", tags]))
        question = _rich_export_html(item.get("question"), images) or "（139 未返回可用题干）"
        title = plain_text(item.get("title"))
        question_text = plain_text(item.get("question"))
        heading = f"第 {index} 题"
        if title and len(title) <= 60 and title[:20] not in question_text[:80]:
            heading += f" · {title}"
        parts = [f'<section class="question-block{(" force-page" if opt["page_break_each"] and index != len(items) else "")}">']
        parts.extend(['<div class="question-heading">', f'<h2 class="question-title">{html.escape(heading)}</h2>'])
        if opt["include_meta"] and meta:
            parts.append(f'<div class="meta">{html.escape(meta)}</div>')
        parts.extend(['<h3 class="section-title">题目</h3>', '</div>', f'<div class="content">{question}</div>'])
        parts.extend('<div class="answer-line"></div>' for _ in range(opt["blank_lines"]))
        if opt["include_my_answer"] and plain_text(item.get("my_answer")):
            parts.extend(['<h3 class="section-title">原作答</h3>', f'<div class="content">{_rich_export_html(item.get("my_answer"), images)}</div>'])
        if opt["answer_mode"] == "inline":
            parts.append('<div class="answer-box"><h3 class="section-title">答案</h3>')
            parts.append(f'<div class="content">{_rich_export_html(item.get("answer"), images) or "（暂无答案）"}</div>')
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                parts.extend(['<h3 class="section-title">解析</h3>', f'<div class="content">{_rich_export_html(item.get("analysis"), images)}</div>'])
            parts.append('</div>')
        if opt["include_note"] and plain_text(item.get("note")):
            parts.extend(['<h3 class="section-title">复盘笔记</h3>', f'<div class="content">{_rich_export_html(item.get("note"), images)}</div>'])
        parts.append('</section>')
        if index != len(items) and not opt["page_break_each"]:
            parts.append('<div class="divider"></div>')
        sections.append("".join(parts))

    answers = ""
    if opt["answer_mode"] == "separate":
        answer_parts = ['<section class="page-break"><h1 class="document-title">参考答案与解析</h1>']
        for index, item in enumerate(items, 1):
            images = item.get("_remote_images") or {}
            answer_parts.extend([
                f'<h2 class="question-title">第 {index} 题</h2>',
                f'<div class="content">{_rich_export_html(item.get("answer"), images) or "（暂无答案）"}</div>',
            ])
            if opt["include_analysis"] and plain_text(item.get("analysis")):
                answer_parts.extend(['<h3 class="section-title">解析</h3>', f'<div class="content">{_rich_export_html(item.get("analysis"), images)}</div>'])
            if index != len(items):
                answer_parts.append('<div class="divider"></div>')
        answer_parts.append('</section>')
        answers = "".join(answer_parts)

    katex_root = "file:///usr/share/javascript/katex"
    script = r"""
      const normalizeLatex = value => String(value)
        .replace(/[（]/g, '(').replace(/[）]/g, ')')
        .replace(/[，]/g, ',').replace(/[：]/g, ':')
        .replace(/\\?~\\?frac/g, '\\frac');
      renderMathInElement(document.body, {
        delimiters: [
          {left: '$$', right: '$$', display: true},
          {left: '\\[', right: '\\]', display: true},
          {left: '\\(', right: '\\)', display: false},
          {left: '$', right: '$', display: false}
        ],
        ignoredTags: ['script', 'style', 'pre', 'code'],
        throwOnError: false,
        strict: 'ignore',
        preProcess: normalizeLatex
      });
      document.body.dataset.ready = 'true';
    """
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
      <link rel="stylesheet" href="{katex_root}/katex.min.css"><style>{css}</style></head><body>
      <h1 class="document-title">{html.escape(opt['title'])}</h1>
      {f'<div class="subtitle">{html.escape(opt["subtitle"])}</div>' if opt['subtitle'] else ''}
      <div class="summary">共 {len(items)} 道错题</div>
      {''.join(sections)}{answers}
      <script src="{katex_root}/katex.min.js"></script><script src="{katex_root}/contrib/auto-render.js"></script>
      <script>{script}</script></body></html>"""


def _render_html_pdf(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    chrome = next((path for path in (Path("/usr/bin/google-chrome"), Path("/usr/bin/chromium")) if path.exists()), None)
    if not chrome:
        raise RuntimeError("未安装 Chrome，无法生成富文本 PDF")
    with tempfile.TemporaryDirectory(prefix="mistake-pdf-") as temp_dir:
        source = Path(temp_dir) / "mistakes.html"
        target = Path(temp_dir) / "mistakes.pdf"
        source.write_text(_export_html(items, raw_options), encoding="utf-8")
        result = subprocess.run([
            str(chrome), "--headless=new", "--no-sandbox", "--disable-gpu",
            "--allow-file-access-from-files", "--run-all-compositor-stages-before-draw",
            "--virtual-time-budget=1500", "--no-pdf-header-footer",
            f"--print-to-pdf={target}", source.as_uri(),
        ], capture_output=True, text=True, check=False, timeout=75)
        if result.returncode or not target.exists():
            raise RuntimeError((result.stderr or result.stdout or "Chrome PDF 生成失败")[-800:])
        content = target.read_bytes()
    if not content.startswith(b"%PDF"):
        raise RuntimeError("Chrome 返回的文件不是有效 PDF")
    return content


def _pdf_to_long_png(pdf: bytes) -> bytes:
    with tempfile.TemporaryDirectory(prefix="mistake-pdf-pages-") as temp_dir:
        source = Path(temp_dir) / "layout.pdf"
        prefix = Path(temp_dir) / "page"
        source.write_bytes(pdf)
        result = subprocess.run(
            ["pdftoppm", "-png", "-r", "150", str(source), str(prefix)],
            capture_output=True, text=True, check=False, timeout=90,
        )
        if result.returncode:
            raise RuntimeError((result.stderr or "PDF 转图片失败")[-800:])
        pages = list(Path(temp_dir).glob("page-*.png"))
        pages.sort(key=lambda path: int(re.search(r"(\d+)$", path.stem).group(1)))
        if not pages or len(pages) > 30:
            raise RuntimeError("PDF 页数异常，OCR Word 单次最多 30 页")
        images = [Image.open(path).convert("RGB") for path in pages]
        gap = 18
        width = max(image.width for image in images)
        height = sum(image.height for image in images) + gap * (len(images) - 1)
        combined = Image.new("RGB", (width, height), "white")
        y = 0
        for image in images:
            combined.paste(image, ((width - image.width) // 2, y))
            y += image.height + gap
        output = io.BytesIO()
        combined.save(output, format="PNG", optimize=True, dpi=(150, 150))
        for image in images:
            image.close()
        return output.getvalue()


def build_pdf(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    """使用 Chrome 排印 139 富文本、配图和 LaTeX，失败时保留原生 PDF 兜底。"""
    try:
        return _render_html_pdf(items, raw_options)
    except Exception as exc:
        print(f"[export] Chrome 富文本 PDF 失败，回退 ReportLab：{exc}")
        return _build_pdf_legacy(items, raw_options)


def build_layout_image(items: list[dict[str, Any]], raw_options: dict[str, Any] | None) -> bytes:
    """把与 PDF 相同的排版逐页转成一张长图，再交给 CamScanner OCR。"""
    try:
        return _pdf_to_long_png(_render_html_pdf(items, raw_options))
    except Exception as exc:
        print(f"[export] 富文本长图生成失败，回退 Pillow 排版：{exc}")
        return _build_layout_image_legacy(items, raw_options)
